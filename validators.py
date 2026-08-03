import os
import re
import openpyxl
import pandas as pd
import io

def check_duplicate_headers(file_path):
    """
    Check if the Excel file has duplicate headers.
    Returns list of duplicate header names, or empty list.
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet = wb.active
        headers = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                headers = [str(cell).strip() for cell in row if cell is not None]
                break
        wb.close()
        
        seen = set()
        duplicates = []
        for h in headers:
            if h in seen:
                duplicates.append(h)
            else:
                seen.add(h)
        return duplicates
    except Exception as e:
        return [f"Error reading headers: {str(e)}"]

def get_question_max(col_name):
    """
    Returns max marks for a question column name.
    Section A (q1a-q1j): max 2.0
    Section B (q2-q11): max 10.0
    Otherwise: None (no max limit or not a question column)
    """
    col = col_name.strip().lower()
    
    # Check Section A: q1a to q1j
    if re.match(r'^q1[a-j]$', col):
        return 2.0
    
    # Check Section B: q2 to q11
    m = re.match(r'^q(\d+)$', col)
    if m:
        q_num = int(m.group(1))
        if 2 <= q_num <= 11:
            return 10.0
            
    return None

def validate_excel_file(file_path, upload_type):
    """
    Validates an uploaded Excel file.
    Returns:
        is_valid (bool): True if no errors, False if errors exist
        errors (list of dict): Details of each error
        valid_indices (list of int): List of 0-based row indices that are valid
        invalid_indices (list of int): List of 0-based row indices that are invalid
        rows_found (int): Total number of rows checked (excluding ignored blank rows)
    """
    errors = []
    valid_indices = []
    invalid_indices = []
    
    # Check file extension
    ext = os.path.splitext(file_path)[1].lower()
    if ext != '.xlsx':
        errors.append({
            "row_num": "-",
            "problem": "Invalid File Format",
            "column": "-",
            "value": ext,
            "reason": "File must be an Excel spreadsheet (.xlsx)"
        })
        return False, errors, [], [], 0

    # Load file with pandas
    try:
        df = pd.read_excel(file_path)
    except Exception as e:
        errors.append({
            "row_num": "-",
            "problem": "Read Error",
            "column": "-",
            "value": "-",
            "reason": f"Failed to read file: {str(e)}"
        })
        return False, errors, [], [], 0

    if df.empty or len(df.columns) == 0:
        errors.append({
            "row_num": "-",
            "problem": "Empty File",
            "column": "-",
            "value": "-",
            "reason": "The uploaded file does not contain any data."
        })
        return False, errors, [], [], 0

    # Check duplicate headers
    duplicates = check_duplicate_headers(file_path)
    if duplicates:
        errors.append({
            "row_num": "-",
            "problem": "Duplicate Headers",
            "column": ", ".join(duplicates),
            "value": "-",
            "reason": f"Duplicate headers are not allowed: {', '.join(duplicates)}"
        })
        return False, errors, [], [], 0

    # Trim spaces from column names
    df.columns = [str(c).strip() for c in df.columns]
    
    # Trim spaces from all cell values
    for col in df.columns:
        df[col] = df[col].apply(lambda x: x.strip() if isinstance(x, str) else x)

    # Filter out completely blank rows
    non_blank_rows = []
    for idx, row in df.iterrows():
        is_blank = True
        for val in row:
            if pd.notna(val) and val != '':
                is_blank = False
                break
        if not is_blank:
            non_blank_rows.append(idx)
            
    df_clean = df.loc[non_blank_rows]
    rows_found = len(df_clean)

    if rows_found == 0:
        errors.append({
            "row_num": "-",
            "problem": "Empty File",
            "column": "-",
            "value": "-",
            "reason": "The file contains only blank rows."
        })
        return False, errors, [], [], 0

    # Type-specific validation
    if upload_type == 'internal':
        # Required columns: Roll_Number, Internal_Marks
        if 'Roll_Number' not in df_clean.columns or 'Internal_Marks' not in df_clean.columns:
            errors.append({
                "row_num": "-",
                "problem": "Missing Columns",
                "column": "-",
                "value": "-",
                "reason": "Missing required column(s): Roll_Number or Internal_Marks"
            })
            return False, errors, [], [], rows_found
            
        # Duplicate Roll Number check
        seen_rolls = set()
        duplicate_rolls = set()
        for idx, row in df_clean.iterrows():
            roll = row['Roll_Number']
            if pd.notna(roll) and str(roll).strip() != '':
                r_str = str(roll).strip()
                if r_str in seen_rolls:
                    duplicate_rolls.add(r_str)
                seen_rolls.add(r_str)

        # Load existing roll numbers from database
        from models import Student
        existing_rolls = {s.roll_no for s in Student.query.all()}

        for idx, row in df_clean.iterrows():
            row_errors = []
            row_num = idx + 2  # 1-based index (header is row 1)
            roll = row['Roll_Number']
            marks = row['Internal_Marks']

            # Check Roll Number
            if pd.isna(roll) or str(roll).strip() == '':
                row_errors.append({
                    "row_num": row_num,
                    "problem": "Missing Roll Number",
                    "column": "Roll_Number",
                    "value": "",
                    "reason": "Roll Number must not be empty."
                })
            else:
                r_str = str(roll).strip()
                if r_str in duplicate_rolls:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Duplicate Roll Number",
                        "column": "Roll_Number",
                        "value": r_str,
                        "reason": "Duplicate Roll Numbers are not allowed."
                    })
                elif r_str not in existing_rolls:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Missing Student",
                        "column": "Roll_Number",
                        "value": r_str,
                        "reason": f"Student with Roll Number '{r_str}' does not exist in the Student database."
                    })

            # Check Internal Marks
            if pd.isna(marks) or str(marks).strip() == '':
                row_errors.append({
                    "row_num": row_num,
                    "problem": "Blank Marks",
                    "column": "Internal_Marks",
                    "value": "",
                    "reason": "Internal marks must not be empty."
                })
            else:
                try:
                    m_val = float(marks)
                    if m_val < 0:
                        row_errors.append({
                            "row_num": row_num,
                            "problem": "Negative Marks",
                            "column": "Internal_Marks",
                            "value": str(marks),
                            "reason": f"Internal marks must be between 0 and 30 (got {m_val})."
                        })
                    elif m_val > 30:
                        row_errors.append({
                            "row_num": row_num,
                            "problem": "Internal Marks >30",
                            "column": "Internal_Marks",
                            "value": str(marks),
                            "reason": f"Internal marks must be between 0 and 30 (got {m_val})."
                        })
                except ValueError:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Invalid Marks",
                        "column": "Internal_Marks",
                        "value": str(marks),
                        "reason": f"Internal marks must be numeric (got '{marks}')."
                    })

            if row_errors:
                errors.extend(row_errors)
                invalid_indices.append(idx)
            else:
                valid_indices.append(idx)

    elif upload_type == 'external':
        # Required columns: Unique_ID, and at least one question column
        if 'Unique_ID' not in df_clean.columns:
            errors.append({
                "row_num": "-",
                "problem": "Missing Columns",
                "column": "-",
                "value": "-",
                "reason": "Missing required column: Unique_ID"
            })
            return False, errors, [], [], rows_found

        question_cols = [col for col in df_clean.columns if col.startswith('Q') or col.startswith('q')]
        if not question_cols:
            errors.append({
                "row_num": "-",
                "problem": "Missing Columns",
                "column": "-",
                "value": "-",
                "reason": "No question columns found (must start with 'Q')."
            })
            return False, errors, [], [], rows_found

        # Total columns identification (if exists)
        total_cols = [c for c in df_clean.columns if c.lower() in ('total', 'external_total', 'external total', 'external_marks')]
        total_col = total_cols[0] if total_cols else None

        # Duplicate UID check
        seen_uids = set()
        duplicate_uids = set()
        for idx, row in df_clean.iterrows():
            uid = row['Unique_ID']
            if pd.notna(uid) and str(uid).strip() != '':
                u_str = str(uid).strip()
                if u_str in seen_uids:
                    duplicate_uids.add(u_str)
                seen_uids.add(u_str)

        for idx, row in df_clean.iterrows():
            row_errors = []
            row_num = idx + 2
            uid = row['Unique_ID']

            # Check Unique_ID
            if pd.isna(uid) or str(uid).strip() == '':
                row_errors.append({
                    "row_num": row_num,
                    "problem": "Missing UID",
                    "column": "Unique_ID",
                    "value": "",
                    "reason": "Unique_ID must not be empty."
                })
            else:
                u_str = str(uid).strip()
                if u_str in duplicate_uids:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Duplicate UID",
                        "column": "Unique_ID",
                        "value": u_str,
                        "reason": "Duplicate Unique_IDs are not allowed."
                    })

            # Check question columns
            question_marks = {}
            for q_col in question_cols:
                val = row[q_col]
                if pd.isna(val) or str(val).strip() == '':
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Blank question marks",
                        "column": q_col,
                        "value": "",
                        "reason": f"Marks for question '{q_col}' must not be empty."
                    })
                else:
                    try:
                        q_val = float(val)
                        if q_val < 0:
                            row_errors.append({
                                "row_num": row_num,
                                "problem": "Negative marks",
                                "column": q_col,
                                "value": str(val),
                                "reason": f"Marks for '{q_col}' cannot be negative (got {q_val})."
                            })
                        else:
                            q_max = get_question_max(q_col)
                            if q_max is not None and q_val > q_max:
                                row_errors.append({
                                    "row_num": row_num,
                                    "problem": "Marks exceeding question maximum",
                                    "column": q_col,
                                    "value": str(val),
                                    "reason": f"Marks for '{q_col}' exceed the maximum limit of {q_max} (got {q_val})."
                                })
                            question_marks[q_col] = q_val
                    except ValueError:
                        row_errors.append({
                            "row_num": row_num,
                            "problem": "Non-numeric values",
                            "column": q_col,
                            "value": str(val),
                            "reason": f"Marks for '{q_col}' must be numeric (got '{val}')."
                        })

            # Verify total matches if total column exists
            if total_col and not row_errors:
                tot_val = row[total_col]
                if pd.notna(tot_val) and str(tot_val).strip() != '':
                    try:
                        uploaded_tot = float(tot_val)
                        calculated_tot = sum(question_marks.values())
                        if abs(uploaded_tot - calculated_tot) > 0.01:
                            row_errors.append({
                                "row_num": row_num,
                                "problem": "Total Mismatch",
                                "column": total_col,
                                "value": f"Uploaded: {uploaded_tot}, Calculated: {calculated_tot}",
                                "reason": f"Uploaded total ({uploaded_tot}) does not match the sum of question marks ({calculated_tot})."
                            })
                    except ValueError:
                        row_errors.append({
                            "row_num": row_num,
                            "problem": "Non-numeric values",
                            "column": total_col,
                            "value": str(tot_val),
                            "reason": f"Total column value must be numeric (got '{tot_val}')."
                        })

            if row_errors:
                errors.extend(row_errors)
                invalid_indices.append(idx)
            else:
                valid_indices.append(idx)

    elif upload_type == 'mapping':
        # Required columns: Unique_ID, Roll_Number
        if 'Unique_ID' not in df_clean.columns or 'Roll_Number' not in df_clean.columns:
            errors.append({
                "row_num": "-",
                "problem": "Missing Columns",
                "column": "-",
                "value": "-",
                "reason": "Missing required column(s): Unique_ID or Roll_Number"
            })
            return False, errors, [], [], rows_found

        # Duplicate UID and Roll Number check
        seen_uids = set()
        duplicate_uids = set()
        seen_rolls = set()
        duplicate_rolls = set()
        for idx, row in df_clean.iterrows():
            uid = row['Unique_ID']
            roll = row['Roll_Number']
            if pd.notna(uid) and str(uid).strip() != '':
                u_str = str(uid).strip()
                if u_str in seen_uids:
                    duplicate_uids.add(u_str)
                seen_uids.add(u_str)
            if pd.notna(roll) and str(roll).strip() != '':
                r_str = str(roll).strip()
                if r_str in seen_rolls:
                    duplicate_rolls.add(r_str)
                seen_rolls.add(r_str)

        # Load existing roll numbers from database
        from models import Student
        existing_rolls = {s.roll_no for s in Student.query.all()}

        for idx, row in df_clean.iterrows():
            row_errors = []
            row_num = idx + 2
            uid = row['Unique_ID']
            roll = row['Roll_Number']

            # Check Unique_ID
            if pd.isna(uid) or str(uid).strip() == '':
                row_errors.append({
                    "row_num": row_num,
                    "problem": "Blank UID",
                    "column": "Unique_ID",
                    "value": "",
                    "reason": "Unique_ID must not be empty."
                })
            else:
                u_str = str(uid).strip()
                if u_str in duplicate_uids:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Duplicate UID",
                        "column": "Unique_ID",
                        "value": u_str,
                        "reason": "Duplicate Unique_IDs are not allowed."
                    })

            # Check Roll Number
            if pd.isna(roll) or str(roll).strip() == '':
                row_errors.append({
                    "row_num": row_num,
                    "problem": "Blank Roll Number",
                    "column": "Roll_Number",
                    "value": "",
                    "reason": "Roll Number must not be empty."
                })
            else:
                r_str = str(roll).strip()
                if r_str in duplicate_rolls:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Duplicate Roll Number",
                        "column": "Roll_Number",
                        "value": r_str,
                        "reason": "Duplicate Roll Numbers are not allowed."
                    })
                elif r_str not in existing_rolls:
                    row_errors.append({
                        "row_num": row_num,
                        "problem": "Missing Student",
                        "column": "Roll_Number",
                        "value": r_str,
                        "reason": f"Student with Roll Number '{r_str}' does not exist in the Student database."
                    })

            if row_errors:
                errors.extend(row_errors)
                invalid_indices.append(idx)
            else:
                valid_indices.append(idx)

    is_valid = len(errors) == 0
    return is_valid, errors, valid_indices, invalid_indices, rows_found

def generate_error_report(errors, output_path):
    """
    Generates an Excel spreadsheet containing the validation errors.
    """
    df = pd.DataFrame(errors)
    df = df[['row_num', 'problem', 'column', 'value', 'reason']]
    df.columns = ['Row Number', 'Problem', 'Column', 'Value', 'Reason']
    
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Validation Errors')
        
        workbook = writer.book
        worksheet = writer.sheets['Validation Errors']
        
        # Auto-adjust column widths
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
