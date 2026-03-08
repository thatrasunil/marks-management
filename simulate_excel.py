import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import random
import os

def calculate_grade(total):
    if total >= 90: return 'A+', 10
    elif total >= 80: return 'A', 9
    elif total >= 70: return 'B', 8
    elif total >= 60: return 'C', 7
    elif total >= 50: return 'D', 6
    elif total >= 40: return 'E', 5
    else: return 'F', 0

def simulate_excel():
    input_file = r"d:\My_Projects\Marks Management System\data.txt"
    output_file = r"d:\My_Projects\Marks Management System\marks-management\student_results_simulated.xlsx"
    
    if not os.path.exists(input_file):
        print(f"Error: {input_file} not found.")
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Student Results"
    
    # 1. Main Title Row
    sheet.append(["Student Results Report (Simulated)"])
    sheet.merge_cells('A1:H1')
    title_cell = sheet['A1']
    title_cell.font = Font(size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 30
    
    # 2. Add empty row
    sheet.append([])
    
    # 3. Header Row
    headers = ["Name", "Roll Number", "Internal", "External", "Total", "Grade", "SGPA", "CGPA"]
    sheet.append(headers)
    
    thin_border = Border(left=Side(style='thin', color='D1D5DB'), 
                         right=Side(style='thin', color='D1D5DB'), 
                         top=Side(style='thin', color='D1D5DB'), 
                         bottom=Side(style='thin', color='D1D5DB'))
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    
    for cell in sheet[3]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    light_gray_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    current_row = 4
    with open(input_file, 'r') as f:
        lines = f.readlines()[1:] # Skip header
        
        for line in lines:
            line = line.strip()
            if not line or line == ".": continue
            
            parts = line.split('\t')
            if len(parts) < 6: continue
            
            name = parts[0]
            roll_no = parts[5]
            
            # Simulate marks
            internal = random.uniform(20, 30)
            external = random.uniform(40, 70)
            total = internal + external
            grade, _ = calculate_grade(total)
            
            # Simulated GPAs
            sgpa = round(random.uniform(7.0, 9.5), 2)
            cgpa = round(sgpa - random.uniform(0, 0.5), 2)
            
            row_data = [
                name,
                roll_no,
                round(internal, 2),
                round(external, 2),
                round(total, 2),
                grade,
                sgpa,
                cgpa
            ]
            sheet.append(row_data)
            
            # Styling
            fill_color = light_gray_fill if (current_row % 2 == 0) else white_fill
            for cell in sheet[current_row]:
                cell.fill = fill_color
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            current_row += 1

    # Auto-adjust column widths
    from openpyxl.utils import get_column_letter
    for col_idx in range(1, sheet.max_column + 1):
        column = get_column_letter(col_idx)
        max_length = 0
        for row_idx in range(1, sheet.max_row + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        sheet.column_dimensions[column].width = max_length + 2

    workbook.save(output_file)
    print(f"Simulated Excel report generated successfully: {output_file}")

if __name__ == "__main__":
    simulate_excel()
